from datetime import date, timedelta
from io import BytesIO
from zipfile import ZipFile

from django.contrib.admin.sites import AdminSite
from django.test import RequestFactory, TestCase
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import User, UserProfile

from .admin import AgeGroupFilter, CompetitionEventFilter, GenderFilter, RegistrationAdmin
from .exports import (
    _safe_excel_text,
    competition_report_definitions,
    competition_reports_zip_response,
    registrations_excel_response,
)
from .models import CompetitionEvent, Festival, Pool, Registration


class RegistrationAdminExportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        now = timezone.now()
        festival = Festival.objects.create(
            title="Test festival",
            registration_starts_at=now - timedelta(days=1),
            registration_ends_at=now + timedelta(days=1),
        )
        pool = Pool.objects.create(name="Test pool")
        event_12 = CompetitionEvent.objects.create(
            festival=festival, title="Event 12", min_age=12, max_age=13
        )
        event_16 = CompetitionEvent.objects.create(
            festival=festival, title="Event 16", min_age=16, max_age=18
        )

        male = User.objects.create_user(
            phone="09120000001", first_name="Ali", last_name="Test"
        )
        UserProfile.objects.create(
            user=male,
            national_id="0012345678",
            birth_date=date(2014, 1, 1),
            gender=UserProfile.Gender.MALE,
        )
        female = User.objects.create_user(
            phone="09120000002", first_name="Sara", last_name="Test"
        )
        UserProfile.objects.create(
            user=female,
            national_id="0012345679",
            birth_date=date(2010, 1, 1),
            gender=UserProfile.Gender.FEMALE,
        )

        common = {
            "festival": festival,
            "pool": pool,
            "insurance_document": "test/insurance.jpg",
            "portrait": "test/portrait.jpg",
            "birth_certificate": "test/birth-certificate.jpg",
            "payment_receipt": "test/payment.jpg",
            "terms_accepted_at": now,
            "insurance_confirmed_at": now,
            "payment_confirmed_at": now,
        }
        male_registration = Registration.objects.create(
            user=male, age_at_registration=12, **common
        )
        male_registration.events.add(event_12)
        female_registration = Registration.objects.create(
            user=female, age_at_registration=16, **common
        )
        female_registration.events.add(event_16)

        cls.event_12 = event_12
        cls.admin_user = User.objects.create_superuser(
            phone="09120000003",
            password="test-admin-password",
            first_name="Admin",
            last_name="Test",
        )

    def setUp(self):
        self.request = RequestFactory().get("/admin/festival/registration/")
        self.model_admin = RegistrationAdmin(Registration, AdminSite())

    def test_registration_filters(self):
        queryset = Registration.objects.all()

        gender_filter = GenderFilter(
            self.request, {"gender": ["female"]}, Registration, self.model_admin
        )
        age_filter = AgeGroupFilter(
            self.request, {"age_group": ["12-13"]}, Registration, self.model_admin
        )
        event_filter = CompetitionEventFilter(
            self.request,
            {"competition_event": [str(self.event_12.pk)]},
            Registration,
            self.model_admin,
        )

        self.assertEqual(gender_filter.queryset(self.request, queryset).count(), 1)
        self.assertEqual(age_filter.queryset(self.request, queryset).count(), 1)
        self.assertEqual(event_filter.queryset(self.request, queryset).count(), 1)

    def test_excel_response_is_valid_and_preserves_identifiers(self):
        queryset = Registration.objects.select_related(
            "user", "user__profile", "festival", "pool"
        ).prefetch_related("events")
        response = registrations_excel_response(queryset)
        workbook = load_workbook(BytesIO(response.content))
        worksheet = workbook["ثبت‌نام‌ها"]

        self.assertEqual(response.status_code, 200)
        self.assertEqual(worksheet.freeze_panes, "A5")
        self.assertEqual(worksheet["A1"].value, "ثبت‌نام‌های جشنواره ندای امید")
        self.assertEqual(worksheet["N4"].value, "تاریخ ثبت‌نام")
        self.assertNotIn("وضعیت", [cell.value for cell in worksheet[4]])
        self.assertEqual(worksheet["E5"].value, "0012345678")
        self.assertEqual(worksheet["E5"].number_format, "@")
        self.assertRegex(worksheet["H5"].value, r"^۱۳\d{2}/\d{2}/\d{2}$")
        self.assertIn("RegistrationData", worksheet.tables)

    def test_formula_injection_is_escaped(self):
        self.assertEqual(_safe_excel_text("=HYPERLINK('bad')"), "'=HYPERLINK('bad')")

    def test_top_export_button_respects_current_filters(self):
        request = RequestFactory().get(
            "/admin/festival/registration/", {"gender": "female"}
        )
        request.user = self.admin_user

        response = self.model_admin.export_filtered_results(request)
        workbook = load_workbook(BytesIO(response.content))
        worksheet = workbook["ثبت‌نام‌ها"]

        self.assertEqual(worksheet.max_row - 4, 1)
        self.assertEqual(worksheet["C5"].value, "Sara")

    def test_competition_reports_zip_has_gender_and_event_reports(self):
        events = CompetitionEvent.objects.order_by("order", "id")
        response = competition_reports_zip_response(Registration.objects.all(), events)
        expected_event_reports = len(competition_report_definitions(events)) * 2

        with ZipFile(BytesIO(response.content)) as archive:
            filenames = archive.namelist()
            self.assertEqual(len(filenames), expected_event_reports + 2)
            self.assertIn("همه-ثبت‌نام‌های-آقایان.xlsx", filenames)
            self.assertIn("همه-ثبت‌نام‌های-خانم‌ها.xlsx", filenames)
            self.assertTrue(all(name.endswith(".xlsx") for name in filenames))
            workbook = load_workbook(BytesIO(archive.read(filenames[0])))
            self.assertEqual(workbook["ثبت‌نام‌ها"].freeze_panes, "A5")
