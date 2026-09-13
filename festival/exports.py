from io import BytesIO
from urllib.parse import quote
from zipfile import ZIP_DEFLATED, ZipFile

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADER_FILL = PatternFill("solid", fgColor="0B4F8A")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(color="0B2D4D", bold=True, size=16)
TEXT_FORMAT = "@"
PERSIAN_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")


def _safe_excel_text(value):
    """Prevent user-controlled text from being interpreted as an Excel formula."""
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def _safe_filename_text(value):
    """Keep archive entries readable without allowing path-like file names."""
    return "".join(
        "-" if character in '\\\\/:*?\"<>|' else character
        for character in str(value)
    ).strip(" .")


def _persian_digits(value):
    return str(value).translate(PERSIAN_DIGITS)


def _gregorian_to_jalali(value):
    gy = value.year - 1600
    gm = value.month - 1
    gd = value.day - 1
    gregorian_days = 365 * gy + (gy + 3) // 4 - (gy + 99) // 100 + (gy + 399) // 400
    month_lengths = (31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31)
    gregorian_days += sum(month_lengths[:gm])
    if gm > 1 and (gy % 4 == 0 and (gy % 100 != 0 or gy % 400 == 0)):
        gregorian_days += 1
    gregorian_days += gd

    jalali_days = gregorian_days - 79
    jalali_cycles, jalali_days = divmod(jalali_days, 12053)
    jy = 979 + 33 * jalali_cycles + 4 * (jalali_days // 1461)
    jalali_days %= 1461
    if jalali_days >= 366:
        jy += (jalali_days - 1) // 365
        jalali_days = (jalali_days - 1) % 365
    if jalali_days < 186:
        jm = 1 + jalali_days // 31
        jd = 1 + jalali_days % 31
    else:
        jm = 7 + (jalali_days - 186) // 30
        jd = 1 + (jalali_days - 186) % 30
    return jy, jm, jd


def _jalali_date(value):
    year, month, day = _gregorian_to_jalali(value)
    return _persian_digits(f"{year:04d}/{month:02d}/{day:02d}")


def _jalali_datetime(value):
    localized = timezone.localtime(value)
    return f"{_jalali_date(localized.date())}، ساعت {_persian_digits(localized.strftime('%H:%M'))}"


def _jalali_filename_date(value):
    year, month, day = _gregorian_to_jalali(value)
    return _persian_digits(f"{year:04d}-{month:02d}-{day:02d}")


def _attachment_header(filename):
    return f"attachment; filename*=UTF-8''{quote(filename)}"


def _age_group(age):
    if 12 <= age <= 13:
        return "۱۲–۱۳ سال"
    if 14 <= age <= 15:
        return "۱۴–۱۵ سال"
    if 16 <= age <= 18:
        return "۱۶–۱۸ سال"
    return "خارج از رده‌های سنی"


def build_registrations_workbook(queryset, report_title="ثبت‌نام‌های جشنواره ندای امید"):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "ثبت‌نام‌ها"

    registrations = list(queryset)
    generated_at = timezone.localtime()

    worksheet["A1"] = report_title
    worksheet["A1"].font = TITLE_FONT
    worksheet["A2"] = (
        f"تاریخ تهیه گزارش: {_jalali_datetime(generated_at)} | تعداد ثبت‌نام‌ها: {_persian_digits(len(registrations))}"
    )
    worksheet["A2"].font = Font(color="52677A", italic=True)

    headers = (
        "ردیف",
        "کد پیگیری",
        "نام",
        "نام خانوادگی",
        "کد ملی",
        "شماره همراه",
        "جنسیت",
        "تاریخ تولد",
        "رده سنی",
        "سن",
        "جشنواره",
        "استخر",
        "مواد مسابقه",
        "تاریخ ثبت‌نام",
    )
    worksheet.append([])
    worksheet.append(headers)

    for cell in worksheet[4]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_number, registration in enumerate(registrations, start=1):
        profile = registration.user.profile
        event_titles = ", ".join(
            _safe_excel_text(event.title) for event in registration.events.all()
        )
        worksheet.append(
            (
                _persian_digits(row_number),
                _safe_excel_text(registration.tracking_code),
                _safe_excel_text(registration.user.first_name),
                _safe_excel_text(registration.user.last_name),
                _safe_excel_text(profile.national_id),
                _safe_excel_text(registration.user.phone),
                "آقا" if profile.gender == "male" else "خانم",
                _jalali_date(profile.birth_date),
                _age_group(registration.age_at_registration),
                _persian_digits(registration.age_at_registration),
                _safe_excel_text(registration.festival.title),
                _safe_excel_text(registration.pool.name),
                event_titles,
                _jalali_datetime(registration.created_at),
            )
        )

    last_row = max(4, worksheet.max_row)
    if registrations:
        table = Table(displayName="RegistrationData", ref=f"A4:N{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)
    else:
        worksheet.auto_filter.ref = "A4:N4"

    worksheet.freeze_panes = "A5"
    worksheet.row_dimensions[4].height = 26
    worksheet.column_dimensions["A"].width = 8
    worksheet.column_dimensions["B"].width = 22
    worksheet.column_dimensions["C"].width = 18
    worksheet.column_dimensions["D"].width = 22
    worksheet.column_dimensions["E"].width = 16
    worksheet.column_dimensions["F"].width = 16
    worksheet.column_dimensions["G"].width = 12
    worksheet.column_dimensions["H"].width = 15
    worksheet.column_dimensions["I"].width = 17
    worksheet.column_dimensions["J"].width = 9
    worksheet.column_dimensions["K"].width = 30
    worksheet.column_dimensions["L"].width = 22
    worksheet.column_dimensions["M"].width = 48
    worksheet.column_dimensions["N"].width = 25

    for row in worksheet.iter_rows(min_row=5, max_row=last_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[4].number_format = TEXT_FORMAT
        row[5].number_format = TEXT_FORMAT
        row[7].number_format = TEXT_FORMAT
        row[13].number_format = TEXT_FORMAT

    workbook.properties.creator = "سامانه جشنواره ندای امید"
    workbook.properties.title = report_title
    return workbook


def registrations_excel_bytes(queryset, report_title="ثبت‌نام‌های جشنواره ندای امید"):
    workbook = build_registrations_workbook(queryset, report_title)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def registrations_excel_response(queryset):
    output = registrations_excel_bytes(queryset)

    filename = f"ثبت‌نام‌های-جشنواره-{_jalali_filename_date(timezone.localdate())}.xlsx"
    response = HttpResponse(
        output,
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = _attachment_header(filename)
    response["X-Content-Type-Options"] = "nosniff"
    return response


AGE_GROUPS = (
    (12, 13, "۱۲–۱۳ سال"),
    (14, 15, "۱۴–۱۵ سال"),
    (16, 18, "۱۶–۱۸ سال"),
)


def competition_report_definitions(events):
    """Split events by official age groups for the 10 competition reports."""
    reports = []
    for event in events:
        for minimum_age, maximum_age, age_label in AGE_GROUPS:
            if event.min_age <= maximum_age and event.max_age >= minimum_age:
                reports.append((event, minimum_age, maximum_age, age_label))
    return reports


def competition_reports_zip_response(registration_queryset, events):
    """Build 20 event-by-gender reports plus two gender summary reports."""
    base_queryset = registration_queryset.select_related(
        "user", "user__profile", "festival", "pool"
    ).prefetch_related("events")
    report_definitions = competition_report_definitions(events)
    output = BytesIO()

    with ZipFile(output, "w", compression=ZIP_DEFLATED, compresslevel=6) as archive:
        for gender, gender_label in (("male", "آقایان"), ("female", "خانم‌ها")):
            for sequence, (event, minimum_age, maximum_age, age_label) in enumerate(
                report_definitions, start=1
            ):
                queryset = base_queryset.filter(
                    user__profile__gender=gender,
                    events=event,
                    age_at_registration__range=(minimum_age, maximum_age),
                ).distinct()
                report_title = (
                    f"{gender_label} | {event.title} | رده سنی {age_label}"
                )
                filename = (
                    f"{gender_label}-{_persian_digits(f'{sequence:02d}')}-"
                    f"{_safe_filename_text(event.title)}-رده-{age_label}.xlsx"
                )
                archive.writestr(
                    filename,
                    registrations_excel_bytes(queryset, report_title),
                )

        for gender, gender_label in (("male", "آقایان"), ("female", "خانم‌ها")):
            queryset = base_queryset.filter(user__profile__gender=gender)
            archive.writestr(
                f"همه-ثبت‌نام‌های-{gender_label}.xlsx",
                registrations_excel_bytes(
                    queryset, f"همه ثبت‌نام‌های {gender_label}"
                ),
            )

    response = HttpResponse(output.getvalue(), content_type="application/zip")
    response["Content-Disposition"] = _attachment_header(
        f"گزارش‌های-مسابقات-{_jalali_filename_date(timezone.localdate())}.zip"
    )
    response["X-Content-Type-Options"] = "nosniff"
    return response
